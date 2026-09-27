"""Lettura e controllo dei file che carica l'utente: TUTTO quello che arriva dall'utente passa da qui.

Ogni file viene:
  1. aperto (se non è un Excel leggibile: Err 30);
  2. cercata l'intestazione nelle prime righe e riconosciute le colonne anche con nomi diversi
     (maiuscole, accenti, punti, spazi non contano);
  3. controllato: colonne obbligatorie, righe vuote, tessere ripetute;
  4. restituito con i nomi di colonna che si aspetta il resto del programma.
Se c'è un problema: ErroreUtente con codice (vedi codici_errori.csv) -> l'app mostra il messaggio
e "cosa fare" e riporta al modulo, senza saltare.

  leggi_iscritti(file, doppio)   elenco iscritti SGAT dei tornei (tennis / beach)
  leggi_sgat(file)               file SGAT del campionato (fogli Casa e Ospite)
  modello_sgat()                 il file Excel "modello" del campionato da scaricare dalla pagina
  modulo_circoli(circoli, prov)  il file Excel dei circoli non trovati, da compilare (provincia)
  leggi_circoli_compilati(file)  lo stesso file compilato dall'utente -> tabella Circolo / Provincia
"""
from __future__ import annotations

import io
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core import normalize
from core.errori import ErroreUtente

RIGHE_DOVE_CERCARE_INTESTAZIONE = 10


def _log(*_a, **_k):
    pass


def _norm(testo) -> str:
    """'Num. Tessera ' -> 'num. tessera' ; 'Cat. Età' -> 'cat. eta' (per confrontare i nomi)."""
    t = unicodedata.normalize("NFKD", str(testo if testo is not None else "")).encode("ascii", "ignore").decode()
    return " ".join(t.lower().split())


def _vuota(valore) -> bool:
    return valore is None or (isinstance(valore, float) and pd.isna(valore)) or str(valore).strip() == ""


def _apri_excel(percorso: Path) -> dict[str, pd.DataFrame]:
    """Tutti i fogli, senza intestazione (la cerchiamo noi). Err 30 se non è un Excel leggibile."""
    try:
        return pd.read_excel(percorso, sheet_name=None, header=None, dtype=object)
    except ImportError:
        raise                                    # libreria mancante = problema di installazione, non dell'utente
    except Exception as e:
        raise ErroreUtente(f"Err 30: «{Path(percorso).name}» non è un file Excel leggibile "
                           f"({type(e).__name__}).") from e


def _riconosci(riga: list, colonne: dict[str, tuple[str, ...]]) -> dict[str, int]:
    """Nome ufficiale -> indice di colonna, per le celle della riga che corrispondono a un alias."""
    trovate: dict[str, int] = {}
    for indice, cella in enumerate(riga):
        testo = _norm(cella)
        for ufficiale, alias in colonne.items():
            if ufficiale not in trovate and testo in alias:
                trovate[ufficiale] = indice
                break
    return trovate


def _tabella(grezzo: pd.DataFrame, colonne: dict, obbligatorie: tuple, dove: str, codice: str) -> pd.DataFrame:
    """Cerca l'intestazione nelle prime righe e restituisce i dati con i nomi ufficiali delle colonne."""
    migliore: dict[str, int] = {}
    for i in range(min(RIGHE_DOVE_CERCARE_INTESTAZIONE, len(grezzo))):
        trovate = _riconosci(grezzo.iloc[i].tolist(), colonne)
        if all(c in trovate for c in obbligatorie):
            dati = grezzo.iloc[i + 1:]
            # le colonne riconosciute prendono il nome ufficiale, le altre restano come sono
            nomi = {j: str(v).strip() for j, v in enumerate(grezzo.iloc[i].tolist()) if not _vuota(v)}
            nomi.update({j: ufficiale for ufficiale, j in trovate.items()})
            tabella = dati[[j for j in nomi]].copy()
            tabella.columns = [nomi[j] for j in nomi]
            if len(tabella):                     # (con zero righe il filtro toglierebbe anche le colonne)
                vuote = [all(_vuota(v) for v in riga) for riga in tabella.itertuples(index=False)]
                tabella = tabella.loc[[not v for v in vuote]]
            return tabella.reset_index(drop=True)
        if len(trovate) > len(migliore):
            migliore = trovate
    mancano = [c for c in obbligatorie if c not in migliore]
    elenco = ", ".join(f"«{n}»" for n in mancano)
    raise ErroreUtente(f"{codice}: in {dove} non trovo le colonne {elenco} (servono: {', '.join(obbligatorie)}).")


# =============================================================================
# Elenco iscritti dei tornei (tennis / beach)
# =============================================================================
COLONNE_ISCRITTI = {                     # nome ufficiale -> come può chiamarsi (già normalizzato)
    "Cognome": ("cognome",),
    "Nome": ("nome",),
    "Num. Tessera": ("num. tessera", "n. tessera", "numero tessera", "num tessera", "tessera", "n.tessera"),
    "Circolo": ("circolo", "societa", "affiliato"),
    "Class.": ("class.", "class", "classifica"),
    "Cat. Età": ("cat. eta", "cat.eta", "cat eta", "categoria eta", "categoria"),
    "Telefono": ("telefono", "cellulare", "tel.", "tel"),
    "Email": ("email", "e-mail", "mail"),
    "Numero Squadra": ("numero squadra", "num. squadra", "n. squadra", "squadra"),
}
OBBLIGATORIE_ISCRITTI = ("Cognome", "Nome", "Num. Tessera", "Circolo", "Class.")


def leggi_iscritti(percorso, doppio: bool = False, log=None) -> pd.DataFrame:
    """Elenco iscritti SGAT -> tabella con i nomi di colonna usati da players/contatti/live_score."""
    log = log or _log
    nome_file = Path(percorso).name
    fogli = _apri_excel(percorso)
    primo = next(iter(fogli.values()))               # l'export SGAT ha un solo foglio: si usa il primo
    df = _tabella(primo, COLONNE_ISCRITTI, OBBLIGATORIE_ISCRITTI, f"«{nome_file}»", "Err 37")
    df = df[~(df["Cognome"].map(_vuota) & df["Nome"].map(_vuota))].reset_index(drop=True)
    if df.empty:
        raise ErroreUtente(f"Err 38: «{nome_file}» non contiene nessun iscritto.")
    if doppio and "Numero Squadra" not in df.columns:
        raise ErroreUtente(f"Err 36: hai scelto Doppio ma in «{nome_file}» non c'è la colonna «Numero Squadra»: "
                           f"serve l'elenco iscritti del doppio esportato da SGAT.")
    return df


# =============================================================================
# File SGAT del campionato (fogli Casa e Ospite)
# =============================================================================
COLONNE_SGAT = {                         # nomi ufficiali minuscoli: li legge players.campionato
    "tessera": ("tessera", "num. tessera", "n. tessera", "numero tessera", "num tessera", "n.tessera"),
    "cognome": ("cognome",),
    "nome": ("nome",),
    "class": ("class", "class.", "classifica"),
}
OBBLIGATORIE_SGAT = ("tessera", "cognome", "nome")
FOGLI_SGAT = ("casa", "ospite")


def _foglio(fogli: dict, ruolo: str, nome_file: str) -> str:
    """Nome vero del foglio Casa / Ospite (maiuscole e spazi non contano)."""
    esatti = [n for n in fogli if _norm(n) == ruolo]
    simili = [n for n in fogli if ruolo[:5] in _norm(n)]           # "casa", "ospit"
    candidati = esatti or simili
    if len(candidati) != 1:
        presenti = ", ".join(f"«{n}»" for n in fogli)
        problema = "manca il foglio" if not candidati else "ci sono più fogli possibili per"
        raise ErroreUtente(f"Err 10: in «{nome_file}» {problema} «{ruolo.capitalize()}» (fogli presenti: {presenti}).")
    return candidati[0]


def leggi_sgat(percorso, log=None, ruoli=FOGLI_SGAT) -> dict[str, pd.DataFrame]:
    """File SGAT del campionato -> {"casa": tabella, "ospite": tabella} con colonne tessera/cognome/nome/class.

    `ruoli` permette di leggere un foglio solo (foglio_cs lo chiama una volta per squadra)."""
    log = log or _log
    nome_file = Path(percorso).name
    fogli = _apri_excel(percorso)
    risultato = {}
    for ruolo in ruoli:
        nome_foglio = _foglio(fogli, ruolo, nome_file)
        df = _tabella(fogli[nome_foglio], COLONNE_SGAT, OBBLIGATORIE_SGAT, f"foglio «{nome_foglio}»", "Err 31")
        df = df[["tessera", "cognome", "nome"] + (["class"] if "class" in df.columns else [])].copy()
        if "class" not in df.columns:
            df["class"] = ""
            log(f"Att. 35: nel foglio «{nome_foglio}» manca la colonna Classifica: nel confronto la classifica "
                f"SGAT resterà vuota.")
        df["tessera"] = df["tessera"].map(normalize.normalizza_tessera)
        senza = df[df["tessera"] == ""]
        if not senza.empty:
            nomi = ", ".join(f"{c} {n}".strip() for c, n in zip(senza["cognome"].fillna(""), senza["nome"].fillna("")))
            log(f"Att. 32: foglio «{nome_foglio}»: {len(senza)} righe senza numero di tessera, escluse: {nomi}")
            df = df[df["tessera"] != ""]
        doppie = df[df["tessera"].duplicated()]["tessera"].unique().tolist()
        if doppie:
            log(f"Att. 33: foglio «{nome_foglio}»: tessere ripetute {', '.join(doppie)}: tengo solo la prima riga.")
            df = df.drop_duplicates("tessera")
        if df.empty:
            raise ErroreUtente(f"Err 34: il foglio «{nome_foglio}» di «{nome_file}» non contiene giocatori.")
        risultato[ruolo] = df.reset_index(drop=True)
    return risultato


def modello_sgat() -> bytes:
    """Il file Excel modello per il campionato: fogli Casa e Ospite con le colonne giuste + istruzioni."""
    wb = Workbook()
    intestazione = ["Tessera", "Cognome", "Nome", "Classifica"]
    for i, nome in enumerate(("Casa", "Ospite")):
        ws = wb.active if i == 0 else wb.create_sheet(nome)
        ws.title = nome
        ws.append(intestazione)
        for cella in ws[1]:
            cella.font = Font(bold=True, color="FFFFFF")
            cella.fill = PatternFill("solid", fgColor="4F81BD")
        for colonna, larghezza in zip("ABCD", (16, 24, 24, 12)):
            ws.column_dimensions[colonna].width = larghezza
        ws.freeze_panes = "A2"
    istruzioni = wb.create_sheet("Istruzioni")
    for riga in (
        ["Modello del file SGAT per il campionato a squadre (FITP Manager)"],
        [""],
        ["1. Nel foglio «Casa» i giocatori della squadra di casa, nel foglio «Ospite» quelli ospiti."],
        ["2. Una riga per giocatore: numero di tessera, cognome, nome e classifica (es. 2.8)."],
        ["3. Non rinominare i fogli e non cambiare l'intestazione della prima riga."],
        ["4. Questo foglio «Istruzioni» viene ignorato dal programma."],
    ):
        istruzioni.append(riga)
    istruzioni["A1"].font = Font(bold=True)
    istruzioni.column_dimensions["A"].width = 100
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# =============================================================================
# Circoli non trovati: file da compilare e sua rilettura
# =============================================================================
# File "circoli da mappare": provincia da tendina, regione automatica; ricaricato, viene aggiunto al csv
FOGLIO_CIRCOLI = "Circoli da mappare"
COLONNE_CIRCOLI = {"Circolo": ("circolo",), "Provincia": ("provincia",)}


def modulo_circoli(circoli: list[str], province: dict[str, str]) -> bytes:
    """Excel da compilare: una riga per circolo, provincia a tendina, regione automatica.

    `province`: provincia -> regione, scritte come in circolo_regioni.csv (data_update.province_circoli)."""
    wb = Workbook()
    ws = wb.active
    ws.title = FOGLIO_CIRCOLI
    ws.append(["Circolo", "Provincia", "Regione (automatica)"])
    for cella in ws[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="4F81BD")
    elenco = wb.create_sheet("Province")               # elenco per la tendina (foglio nascosto)
    elenco.append(["Provincia", "Regione"])
    for provincia, regione in sorted(province.items()):
        elenco.append([provincia, regione])
    elenco.sheet_state = "hidden"
    ultima = len(province) + 1

    tendina = DataValidation(type="list", formula1=f"Province!$A$2:$A${ultima}", allow_blank=True,
                             showErrorMessage=True, errorTitle="Provincia non valida",
                             error="Scegli la provincia dal menu a tendina.")
    ws.add_data_validation(tendina)
    grigio = PatternFill("solid", fgColor="EEEEEE")
    for riga, circolo in enumerate(circoli, start=2):
        ws.cell(riga, 1, circolo)
        tendina.add(f"B{riga}")
        regione = ws.cell(riga, 3, f'=IF(B{riga}="","",IFERROR(VLOOKUP(B{riga},Province!$A$2:$B${ultima},2,FALSE),"provincia non valida"))')
        regione.fill = grigio
    for colonna, larghezza in zip("ABC", (60, 28, 30)):
        ws.column_dimensions[colonna].width = larghezza
    ws.freeze_panes = "A2"

    istruzioni = wb.create_sheet("Istruzioni", 1)
    for riga in (
        ["Circoli non trovati in data/circolo_regioni.csv (nel foglio Giocatori hanno regione ND)"],
        [""],
        ["1. Nel foglio «Circoli da mappare», per ogni circolo scegli la provincia DEL CIRCOLO dal menu a tendina (colonna B)."],
        ["   La regione (colonna C) compare da sola: serve solo per controllare."],
        ["2. Se non vuoi aggiungere un circolo, lascia vuota la sua provincia."],
        ["3. Non cambiare il nome del circolo: deve restare identico a quello di SGAT."],
        ["4. Salva e ricarica questo file nell'app: pagina finale dell'elaborazione, oppure scheda «Circoli»"],
        ["   nella pagina Tennis o Beach. I circoli vengono aggiunti in fondo a circolo_regioni.csv;"],
        ["   la versione precedente del file resta in data/archivio/."],
    ):
        istruzioni.append(riga)
    istruzioni["A1"].font = Font(bold=True)
    istruzioni.column_dimensions["A"].width = 110
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def leggi_circoli_compilati(percorso) -> pd.DataFrame:
    """File "circoli da mappare" compilato -> tabella con Circolo e Provincia (testi, righe vuote tolte).

    La colonna Regione NON si legge: è una formula (senza Excel non ha valore) e la ricava comunque
    il programma dalla provincia."""
    nome_file = Path(percorso).name
    fogli = _apri_excel(percorso)
    nome_foglio = next((n for n in fogli if _norm(n) == _norm(FOGLIO_CIRCOLI)), next(iter(fogli)))
    df = _tabella(fogli[nome_foglio], COLONNE_CIRCOLI, ("Circolo", "Provincia"), f"«{nome_file}»", "Err 48")
    df = df[["Circolo", "Provincia"]].copy()
    # il nome del circolo resta IDENTICO (solo spazi ai lati): è la chiave usata da players.mappa_clubs
    df["Circolo"] = df["Circolo"].map(lambda v: "" if _vuota(v) else str(v).strip())
    df["Provincia"] = df["Provincia"].map(lambda v: "" if _vuota(v) else " ".join(str(v).split()))
    return df[df["Circolo"] != ""].reset_index(drop=True)

