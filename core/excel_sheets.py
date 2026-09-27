"""Copia di fogli Excel mantenendo la formattazione.

Prima esistevano TRE funzioni quasi identiche:
- copia_file.copy_sheet_with_format()
- file_torneo.aggiungi_file()
- crea_entry.aggiungi_entry_list()

Qui c'e' UNA sola funzione parametrica, `copia_foglio`, piu' una comoda
`copia_fogli` per copiarne diversi in un colpo. La logica di copia cella/stile
e' fattorizzata in `_copia_contenuto`.
"""
from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl


def _copia_contenuto(ws_src, ws_dst) -> None:
    """Copia celle, stili, larghezze colonne, altezze righe e celle unite."""
    for row in ws_src.iter_rows():
        for cell in row:
            new_cell = ws_dst[cell.coordinate]
            new_cell.value = cell.value
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.fill = copy(cell.fill)
                new_cell.border = copy(cell.border)
                new_cell.alignment = copy(cell.alignment)
                new_cell.number_format = cell.number_format
                new_cell.protection = copy(cell.protection)

    for col_letter, col_dim in ws_src.column_dimensions.items():
        ws_dst.column_dimensions[col_letter].width = col_dim.width
    for row_num, row_dim in ws_src.row_dimensions.items():
        ws_dst.row_dimensions[row_num].height = row_dim.height
    for merged_range in ws_src.merged_cells.ranges:
        ws_dst.merge_cells(str(merged_range))


def copia_foglio(source_file, target_file, foglio_sorgente: str | None = None,
                 nuovo_nome: str | None = None, zoom: int | None = None) -> None:
    """Copia UN foglio da `source_file` a `target_file` con la formattazione.

    Args:
        source_file: file Excel di origine.
        target_file: file Excel di destinazione (creato se non esiste).
        foglio_sorgente: nome del foglio da copiare; None = foglio attivo.
        nuovo_nome: nome del foglio nel target; None = stesso nome della sorgente.
        zoom: se indicato, imposta lo zoom del foglio copiato (es. 70).

    Se nel target esiste gia' un foglio con lo stesso nome, viene sostituito.
    """
    wb_src = openpyxl.load_workbook(source_file)
    ws_src = wb_src[foglio_sorgente] if foglio_sorgente else wb_src.active

    if Path(target_file).exists():
        wb_dst = openpyxl.load_workbook(target_file)
    else:
        wb_dst = openpyxl.Workbook()
        wb_dst.remove(wb_dst.active)          # parte vuoto

    nome = nuovo_nome or ws_src.title
    if nome in wb_dst.sheetnames:
        del wb_dst[nome]
    ws_dst = wb_dst.create_sheet(title=nome)

    _copia_contenuto(ws_src, ws_dst)
    if zoom is not None:
        ws_dst.sheet_view.zoomScale = zoom

    wb_dst.save(target_file)


def copia_fogli(source_file, target_file, fogli: list[str]) -> list[str]:
    """Copia PIU' fogli in un solo salvataggio. Restituisce i fogli non trovati.

    Sostituisce `file_torneo.aggiungi_file()`.
    """
    wb_src = openpyxl.load_workbook(source_file)
    wb_dst = openpyxl.load_workbook(target_file)
    non_trovati: list[str] = []

    for nome in fogli:
        if nome not in wb_src.sheetnames:
            non_trovati.append(nome)
            continue
        if nome in wb_dst.sheetnames:
            del wb_dst[nome]
        ws_dst = wb_dst.create_sheet(nome)
        _copia_contenuto(wb_src[nome], ws_dst)

    wb_dst.save(target_file)
    return non_trovati
