"""Stili e formattazione Excel condivisi.

`formatta_file()` era COPIATA IDENTICA in file_torneo.py e crea_entry.py.
Qui esiste una sola versione, piu' i mattoncini di stile riutilizzabili
(header, bordi, evidenziazioni SI/NO, font link) usati anche dal campionato.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import conf as config

# --- Mattoncini di stile riutilizzabili ---------------------------------------
BORDO_SOTTILE = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
CENTRATO = Alignment(horizontal="center", vertical="center")
FONT_GRASSETTO = Font(bold=True)
FONT_LINK = Font(color=config.COLORE_LINK, underline="single")


def fill(colore: str) -> PatternFill:
    """Riempimento solido di un colore esadecimale."""
    return PatternFill(start_color=colore, end_color=colore, fill_type="solid")


def formatta_intestazione(ws, colore_header: str = config.COLORE_HEADER,
                          testo_bianco: bool = False) -> None:
    """Formatta la prima riga come intestazione (fill + grassetto + bordo)."""
    header_fill = fill(colore_header)
    header_font = Font(bold=True, color="FFFFFF") if testo_bianco else FONT_GRASSETTO
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = BORDO_SOTTILE
        cell.alignment = CENTRATO


def formatta_file(ws, colore_header: str = config.COLORE_HEADER,
                  colonna_testo: str | None = "A",
                  larghezza: int = config.LARGHEZZA_COLONNA_DEFAULT) -> None:
    """Formattazione standard di un foglio (ex `formatta_file` duplicata).

    - intestazione in grassetto e colorata
    - `colonna_testo` forzata a formato testo '@' (utile per le tessere)
    - larghezza uniforme delle colonne
    """
    formatta_intestazione(ws, colore_header=colore_header)

    if colonna_testo:
        for cell in ws[colonna_testo][1:]:      # salta l'intestazione
            cell.number_format = "@"

    for col_idx, _ in enumerate(ws.iter_cols(), start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = larghezza


def evidenzia_corrispondenza(ws, nome_colonna: str = "Corrispondenza Dati") -> None:
    """Colora di verde/rosso le celle SI/NO della colonna indicata.

    Usata dal campionato a squadre (ex salva_excel_formattato).
    """
    fill_si = fill(config.COLORE_VERDE)
    fill_no = fill(config.COLORE_ROSSO)

    col_idx = None
    for idx, cell in enumerate(ws[1], 1):
        if cell.value == nome_colonna:
            col_idx = idx
            break
    if col_idx is None:
        return

    for row in ws.iter_rows(min_row=2):
        cella = row[col_idx - 1]
        if cella.value == "SI":
            cella.fill = fill_si
        elif str(cella.value).startswith("NO"):
            cella.fill = fill_no
